from pathlib import Path
import csv

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX.
    Использовать openpyxl ИЛИ xlsxwriter.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """
    in_data_csv = Path(csv_path)
    out_data_xlsx = Path(xlsx_path)

    if not in_data_csv.is_file():
        raise FileNotFoundError(f"Файл не найден: {in_data_csv}")

    if in_data_csv.suffix.lower() != ".csv":
        raise ValueError("Входной файл должен быть формата .csv")

    if out_data_xlsx.suffix.lower() != ".xlsx":
        raise ValueError("Выходной файл должен быть формата .xlsx")
    
    try:
        with in_data_csv.open('r', encoding='utf-8',newline='') as f:
            reader = csv.reader(f, strict=True)
            headers = next(reader, None)

            if not headers:
                raise ValueError("CSV пустой или не содержит заголовков")
            for header in headers:
                if not header.strip():
                    raise ValueError("Заголовки не должны быть пустыми")
            if len(set(headers)) != len(headers):
                raise ValueError("Заголовки не должны повторяться")
            
            data = [headers]

            for record in reader:
                if len(record) == 0:
                    continue
                if len(record) != len(headers):
                    raise ValueError("Количество полей не совпадает с заголовком")
                data.append(record)
    except csv.Error as e:
        raise ValueError("Ошибка в структуре CSV") from e
    
    excel = Workbook()
    worksheet = excel.active
    worksheet.title = 'Sheet1'
    widths = [0] * len(headers)
    
    for record in data:
        worksheet.append(record)
        for index, text in enumerate(record):
            cell = worksheet.cell(row=worksheet.max_row, column=index + 1)
            cell.data_type='s'

            length = max((len(line) for line in text.splitlines()), default = 0)
            if length > widths[index]:
                widths[index]=length
        
        for index, width in enumerate(widths, start=1):
            letter = get_column_letter(index)
            worksheet.column_dimensions[letter].width = min(255, max(8, width+2))
        try:
            excel.save(out_data_xlsx)
        finally:
            excel.close()
